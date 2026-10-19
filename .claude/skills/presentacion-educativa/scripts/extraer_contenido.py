#!/usr/bin/env python3
"""Extrae texto, tablas, notas e imágenes de material original a Markdown.

Uso:
    python3 -I extraer_contenido.py <archivos o carpetas...> --salida <carpeta>

Admite:
    .pptx  (python-pptx)  una sección por diapositiva: textos en orden, tablas,
                          gráficos, SmartArt, notas del orador e imágenes.
    .pdf   (pymupdf)      una sección por página: texto e imágenes.
    .docx  (python-docx)  párrafos (con títulos y viñetas) y tablas en orden.
    imágenes              .png .jpg .jpeg .gif .webp .bmp .tif .tiff .svg:
                          solo se copian y se listan (descríbelas tú).

Escribe <salida>/<nombre>.md por archivo y las imágenes, sin duplicados, en
<salida>/img-originales/. Al final imprime un resumen: diapositivas o páginas
por archivo y, para los pptx, la paleta (colores srgbClr más usados) y las
fuentes. Las rutas de imagen del Markdown son relativas a <salida>.
"""

from __future__ import annotations

import argparse
import hashlib
import shutil
import sys
from collections import Counter
from pathlib import Path

EXT_IMAGEN = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".svg"}
EXT_ADMITIDAS = {".pptx", ".pdf", ".docx"} | EXT_IMAGEN

INSTALAR = {
    "pptx": "pip install python-pptx",
    "pymupdf": "pip install pymupdf",
    "docx": "pip install python-docx",
}
NOMBRE_PAQUETE = {"pptx": "python-pptx", "pymupdf": "pymupdf", "docx": "python-docx"}


class FaltaLibreria(Exception):
    pass


def importar(modulo: str):
    try:
        if modulo == "pymupdf":
            try:
                import pymupdf  # type: ignore
                return pymupdf
            except ImportError:
                import fitz  # type: ignore  (nombre antiguo de PyMuPDF)
                return fitz
        return __import__(modulo)
    except ImportError as e:
        raise FaltaLibreria(
            f"Falta la librería {NOMBRE_PAQUETE[modulo]}. Instálala con:\n    {INSTALAR[modulo]}"
        ) from e


# ---------------------------------------------------------------------------
# Almacén de imágenes sin duplicados
# ---------------------------------------------------------------------------

class Imagenes:
    def __init__(self, carpeta: Path):
        self.carpeta = carpeta
        self.vistas: dict[str, str] = {}   # sha1 -> ruta relativa
        self.nuevas = 0
        self.repetidas = 0

    def guardar(self, datos: bytes, ext: str, nombre: str) -> tuple[str, bool]:
        """Devuelve (ruta relativa a la salida, es_nueva)."""
        h = hashlib.sha1(datos).hexdigest()
        if h in self.vistas:
            self.repetidas += 1
            return self.vistas[h], False
        self.carpeta.mkdir(parents=True, exist_ok=True)
        ext = (ext or "bin").lower().lstrip(".")
        if ext == "jpeg":
            ext = "jpg"
        destino = self.carpeta / f"{nombre}.{ext}"
        n = 2
        while destino.exists():
            destino = self.carpeta / f"{nombre}-{n}.{ext}"
            n += 1
        destino.write_bytes(datos)
        rel = f"{self.carpeta.name}/{destino.name}"
        self.vistas[h] = rel
        self.nuevas += 1
        return rel, True


def dimensiones(datos: bytes) -> str:
    """Tamaño en píxeles si PyMuPDF puede leerlo; si no, cadena vacía."""
    try:
        pm = importar("pymupdf")
        px = pm.Pixmap(datos)
        return f"{px.width}×{px.height} px"
    except Exception:
        return ""


def limpiar(texto: str) -> str:
    return texto.replace("\x0b", "\n").replace("\r", "").strip()


def tabla_md(filas: list[list[str]]) -> list[str]:
    if not filas:
        return []
    ancho = max(len(f) for f in filas)
    filas = [[c.replace("\n", "<br>").replace("|", "\\|") for c in f] + [""] * (ancho - len(f)) for f in filas]
    out = ["| " + " | ".join(filas[0]) + " |", "|" + "---|" * ancho]
    out += ["| " + " | ".join(f) + " |" for f in filas[1:]]
    return out


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
NS_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
NS_DGM = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"


def extraer_pptx(ruta: Path, imgs: Imagenes, stem: str) -> dict:
    pptx = importar("pptx")
    from pptx.enum.shapes import MSO_SHAPE_TYPE  # type: ignore
    from pptx.util import Emu  # type: ignore  # noqa: F401

    prs = pptx.Presentation(str(ruta))
    colores: Counter = Counter()
    fuentes: Counter = Counter()
    lineas = [f"# {ruta.name}  ({len(prs.slides)} diapositivas, "
              f"{prs.slide_width}x{prs.slide_height} EMU)", ""]
    n_img = 0

    # Fuentes del tema (las que se usan cuando el texto no fija ninguna)
    try:
        for master in prs.slide_masters:
            for rel in master.part.rels.values():
                if rel.reltype.endswith("/theme"):
                    from lxml import etree  # type: ignore
                    raiz = etree.fromstring(rel.target_part.blob)
                    for tag, rol in (("majorFont", "títulos"), ("minorFont", "cuerpo")):
                        for f in raiz.iter(f"{NS_A}{tag}"):
                            latin = f.find(f"{NS_A}latin")
                            if latin is not None and latin.get("typeface"):
                                fuentes[f"{latin.get('typeface')} (tema, {rol})"] += 0
    except Exception:
        pass

    def contar_estilo(elemento):
        for c in elemento.iter(f"{NS_A}srgbClr"):
            v = c.get("val")
            if v:
                colores["#" + v.upper()] += 1
        for f in elemento.iter(f"{NS_A}latin"):
            t = f.get("typeface")
            if t and not t.startswith("+"):
                fuentes[t] += 1

    def texto_smartart(shape, slide) -> list[str]:
        out = []
        for rid_el in shape._element.iter(f"{NS_DGM}relIds"):
            rid = rid_el.get(f"{NS_R}dm")
            if not rid:
                continue
            try:
                from lxml import etree  # type: ignore
                parte = slide.part.related_part(rid)
                raiz = etree.fromstring(parte.blob)
                for p in raiz.iter(f"{NS_A}p"):
                    t = "".join(x.text or "" for x in p.iter(f"{NS_A}t")).strip()
                    if t:
                        out.append(t)
            except Exception:
                pass
        return out

    def recorrer(shapes, slide, num, salida: list[str]):
        nonlocal n_img
        for sh in shapes:
            try:
                tipo = sh.shape_type
            except Exception:
                tipo = None
            if tipo == MSO_SHAPE_TYPE.GROUP:
                recorrer(sh.shapes, slide, num, salida)
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                filas = [[limpiar(c.text) for c in fila.cells] for fila in sh.table.rows]
                salida.append("")
                salida += tabla_md(filas)
                salida.append("")
                continue
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                try:
                    titulo = ch.chart_title.text_frame.text if ch.has_title else ""
                except Exception:
                    titulo = ""
                salida.append(f"[GRÁFICO {ch.chart_type}{': ' + titulo if titulo else ''}]")
                try:
                    cats = [str(c) for c in ch.plots[0].categories]
                    salida.append("  Categorías: " + ", ".join(cats))
                    for s in ch.series:
                        salida.append(f"  Serie «{s.name}»: " + ", ".join(str(v) for v in s.values))
                except Exception:
                    pass
                continue
            imagen = None
            if tipo == MSO_SHAPE_TYPE.PICTURE or (getattr(sh, "is_placeholder", False) and hasattr(sh, "image")):
                try:
                    imagen = sh.image
                except Exception:
                    imagen = None
            if imagen is not None:
                n_img += 1
                rel, nueva = imgs.guardar(imagen.blob, imagen.ext, f"{stem}-d{num:02d}-{n_img}")
                dims = dimensiones(imagen.blob)
                alt = ""
                try:
                    alt = sh._element.xpath("./*[1]/p:cNvPr/@descr")
                    alt = alt[0] if alt else ""
                    if Path(alt).suffix.lower() in EXT_IMAGEN:
                        alt = ""   # nombre de archivo, no una descripción
                except Exception:
                    alt = ""
                extra = " · ".join(x for x in (dims, "repetida" if not nueva else "", f"alt: {alt}" if alt else "") if x)
                salida.append(f"[IMAGEN {rel}{' · ' + extra if extra else ''}]")
                continue
            sa = texto_smartart(sh, slide)
            if sa:
                salida.append("[SMARTART]")
                salida += [f"- {t}" for t in sa]
                continue
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = limpiar("".join(r.text for r in p.runs) if p.runs else p.text)
                    if not t:
                        continue
                    nivel = p.level or 0
                    salida.append(("  " * nivel + "- " + t) if nivel else t)

    for num, slide in enumerate(prs.slides, start=1):
        oculta = slide._element.get("show") == "0"
        try:
            layout = slide.slide_layout.name
        except Exception:
            layout = "?"
        lineas.append(f"## Diapositiva {num}  (layout: {layout}){'  [OCULTA]' if oculta else ''}")
        cuerpo: list[str] = []
        recorrer(slide.shapes, slide, num, cuerpo)
        contar_estilo(slide._element)
        lineas += cuerpo
        if slide.has_notes_slide:
            notas = limpiar(slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else "")
            if notas:
                lineas += ["", "**Notas del orador:**", notas]
        lineas.append("")

    # Colores de fondo y de los layouts usados también cuentan para la paleta
    usados = {id(s.slide_layout): s.slide_layout for s in prs.slides}
    for lay in usados.values():
        contar_estilo(lay._element)
    for m in prs.slide_masters:
        contar_estilo(m._element)

    return {"tipo": "pptx", "unidades": len(prs.slides), "etiqueta": "diapositivas",
            "lineas": lineas, "imagenes": n_img, "colores": colores, "fuentes": fuentes}


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def extraer_pdf(ruta: Path, imgs: Imagenes, stem: str) -> dict:
    pm = importar("pymupdf")
    doc = pm.open(str(ruta))
    lineas = [f"# {ruta.name}  ({doc.page_count} páginas)", ""]
    n_img = 0
    sin_texto = 0
    for i, pagina in enumerate(doc, start=1):
        lineas.append(f"## Página {i}")
        texto = limpiar(pagina.get_text("text"))
        if texto:
            lineas.append(texto)
        else:
            sin_texto += 1
            lineas.append("_(sin texto extraíble: probablemente escaneada; mira la imagen)_")
        vistas_pag = set()
        for info in pagina.get_images(full=True):
            xref = info[0]
            if xref in vistas_pag:
                continue
            vistas_pag.add(xref)
            try:
                img = doc.extract_image(xref)
            except Exception:
                continue
            if not img or min(img.get("width", 0), img.get("height", 0)) < 24:
                continue   # filetes, puntos y adornos
            n_img += 1
            rel, nueva = imgs.guardar(img["image"], img.get("ext", "png"), f"{stem}-p{i:02d}-{n_img}")
            lineas.append(f"[IMAGEN {rel} · {img.get('width')}×{img.get('height')} px{'' if nueva else ' · repetida'}]")
        lineas.append("")
    if sin_texto:
        lineas.insert(1, f"> Aviso: {sin_texto} página(s) sin texto extraíble (¿escaneado?). Lee las imágenes.\n")
    return {"tipo": "pdf", "unidades": doc.page_count, "etiqueta": "páginas",
            "lineas": lineas, "imagenes": n_img}


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def extraer_docx(ruta: Path, imgs: Imagenes, stem: str) -> dict:
    docx = importar("docx")
    from docx.table import Table  # type: ignore
    from docx.text.paragraph import Paragraph  # type: ignore

    d = docx.Document(str(ruta))
    lineas = [f"# {ruta.name}", ""]
    n_par = n_tab = 0
    cuerpo = d.element.body
    for hijo in cuerpo.iterchildren():
        etiqueta = hijo.tag.rsplit("}", 1)[-1]
        if etiqueta == "p":
            p = Paragraph(hijo, d)
            t = limpiar(p.text)
            if not t:
                continue
            n_par += 1
            estilo = (p.style.name if p.style is not None else "") or ""
            if estilo.lower().startswith(("heading", "título", "titulo")) and estilo[-1:].isdigit():
                lineas += ["", "#" * min(6, int(estilo[-1]) + 2) + " " + t, ""]
            elif estilo.lower() in ("title", "título", "titulo"):
                lineas += ["", "## " + t, ""]
            elif "list" in estilo.lower() or "lista" in estilo.lower() or hijo.find(
                    ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr") is not None:
                lineas.append("- " + t)
            else:
                lineas += [t, ""]
        elif etiqueta == "tbl":
            n_tab += 1
            tabla = Table(hijo, d)
            filas = [[limpiar(c.text) for c in fila.cells] for fila in tabla.rows]
            lineas += [""] + tabla_md(filas) + [""]

    n_img = 0
    for rel in d.part.rels.values():
        if "image" in rel.reltype:
            try:
                parte = rel.target_part
            except Exception:
                continue
            n_img += 1
            ext = Path(parte.partname).suffix or ".png"
            r, nueva = imgs.guardar(parte.blob, ext, f"{stem}-img-{n_img}")
            dims = dimensiones(parte.blob)
            lineas.append(f"[IMAGEN {r}{' · ' + dims if dims else ''}{'' if nueva else ' · repetida'}]")
    lineas[0] = f"# {ruta.name}  ({n_par} párrafos, {n_tab} tablas, {n_img} imágenes)"
    return {"tipo": "docx", "unidades": n_par, "etiqueta": "párrafos",
            "lineas": lineas, "imagenes": n_img, "tablas": n_tab}


# ---------------------------------------------------------------------------
# Imágenes sueltas
# ---------------------------------------------------------------------------

def extraer_imagen(ruta: Path, imgs: Imagenes, stem: str) -> dict:
    datos = ruta.read_bytes()
    rel, nueva = imgs.guardar(datos, ruta.suffix, stem)
    dims = "" if ruta.suffix.lower() == ".svg" else dimensiones(datos)
    lineas = [f"# {ruta.name}", "",
              f"[IMAGEN {rel}{' · ' + dims if dims else ''}{'' if nueva else ' · repetida'}]", "",
              "_Imagen sin texto extraído: descríbela leyéndola (contenido, textos visibles, colores)._", ""]
    return {"tipo": "imagen", "unidades": 1, "etiqueta": "imagen", "lineas": lineas, "imagenes": 1}


# ---------------------------------------------------------------------------

def hsl(hexa: str) -> tuple[float, float, float]:
    import colorsys
    r, g, b = (int(hexa[i:i + 2], 16) / 255 for i in (1, 3, 5))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    return h * 360, s, l


def es_neutro(hexa: str) -> bool:
    _, s, l = hsl(hexa)
    return s < 0.12 or l > 0.97 or l < 0.04   # blancos, negros y grises


def sugerir_config(colores: Counter) -> list[str]:
    """Propuesta de colores para config a partir de la frecuencia y el HSL."""
    orden = [c for c, _ in colores.most_common(40)]
    sug = []
    medios = [c for c in orden if hsl(c)[1] >= 0.35 and 0.28 <= hsl(c)[2] <= 0.75]
    if medios:
        prim = medios[0]
        sug.append(f"colorPrimario {prim}")
        hp = hsl(prim)[0]
        for c in medios[1:]:
            d = abs(hsl(c)[0] - hp)
            if min(d, 360 - d) > 40:
                sug.append(f"colorAcento {c}")
                break
    oscuros = [c for c in orden if hsl(c)[2] < 0.22 and not es_neutro(c)] or [c for c in orden if hsl(c)[2] < 0.22]
    if oscuros:
        sug.append(f"colorTexto {oscuros[0]}")
    claros = [c for c in orden if 0.88 < hsl(c)[2] < 0.985]
    if claros:
        sug.append(f"colorFondo {claros[0]}")
    return sug


def main() -> int:
    ap = argparse.ArgumentParser(description="Extrae texto e imágenes de pptx, pdf, docx e imágenes a Markdown.")
    ap.add_argument("archivos", nargs="+", help="archivos o carpetas")
    ap.add_argument("--salida", required=True, help="carpeta donde escribir los .md (las imágenes van a img-originales/)")
    a = ap.parse_args()

    salida = Path(a.salida)
    salida.mkdir(parents=True, exist_ok=True)
    imgs = Imagenes(salida / "img-originales")

    entradas: list[Path] = []
    for x in a.archivos:
        p = Path(x)
        if p.is_dir():
            entradas += sorted(q for q in p.iterdir() if q.is_file() and q.suffix.lower() in EXT_ADMITIDAS
                               and not q.name.startswith(("~$", ".")))
        elif p.is_file():
            entradas.append(p)
        else:
            print(f"No existe: {p}", file=sys.stderr)
    if not entradas:
        print("No hay archivos que procesar.", file=sys.stderr)
        return 2

    stems = Counter(p.stem for p in entradas)
    resultados = []
    errores = 0
    faltan: set[str] = set()
    for p in entradas:
        ext = p.suffix.lower()
        stem = p.stem if stems[p.stem] == 1 else f"{p.stem}.{ext.lstrip('.')}"
        try:
            if ext == ".pptx":
                r = extraer_pptx(p, imgs, stem)
            elif ext == ".pdf":
                r = extraer_pdf(p, imgs, stem)
            elif ext == ".docx":
                r = extraer_docx(p, imgs, stem)
            elif ext in EXT_IMAGEN:
                r = extraer_imagen(p, imgs, stem)
            else:
                print(f"Formato no admitido, se salta: {p.name} (admite pptx, pdf, docx e imágenes; "
                      f"un .ppt o .doc antiguo hay que guardarlo antes como pptx/docx)", file=sys.stderr)
                continue
        except FaltaLibreria as e:
            faltan.add(str(e))
            errores += 1
            continue
        except Exception as e:  # archivo dañado o protegido
            print(f"Error al leer {p.name}: {type(e).__name__}: {e}", file=sys.stderr)
            errores += 1
            continue
        md = salida / f"{stem}.md"
        texto = "\n".join(r["lineas"]).rstrip() + "\n"
        while "\n\n\n" in texto:
            texto = texto.replace("\n\n\n", "\n\n")
        md.write_text(texto, encoding="utf-8")
        r["archivo"], r["md"] = p.name, md
        resultados.append(r)

    # ----------------------------- resumen -----------------------------
    if not resultados:
        for f in sorted(faltan):
            print(f"\n{f}", file=sys.stderr)
        print("\nNo se ha extraído nada.", file=sys.stderr)
        return 1
    print(f"\nResumen de la extracción → {salida}")
    for r in resultados:
        extra = f" · {r['tablas']} tablas" if r.get("tablas") else ""
        cuantas = "" if r["tipo"] == "imagen" else f" · {r['imagenes']} imágenes"
        print(f"  {r['archivo']}: {r['unidades']} {r['etiqueta']}{cuantas}{extra} → {r['md'].name}")
    print(f"  Imágenes guardadas en {imgs.carpeta}: {imgs.nuevas} distintas"
          f" ({imgs.repetidas} repeticiones no se han vuelto a guardar)")

    pptxs = [r for r in resultados if r["tipo"] == "pptx"]
    if pptxs:
        colores: Counter = Counter()
        fuentes: Counter = Counter()
        for r in pptxs:
            colores.update(r["colores"])
            fuentes.update(r["fuentes"])
            print(f"\n  Paleta de {r['archivo']}: " +
                  ", ".join(f"{c} ×{n}" for c, n in r["colores"].most_common(8)) or "—")
        if len(pptxs) > 1:
            print("\n  Paleta conjunta (colores srgbClr más usados):")
        else:
            print("\n  Colores srgbClr más usados:")
        for c, n in colores.most_common(12):
            print(f"    {c}  ×{n}{'   (neutro)' if es_neutro(c) else ''}")
        print("  Fuentes:")
        for f, n in sorted(fuentes.items(), key=lambda x: -x[1])[:8]:
            print(f"    {f}" + (f"  ×{n}" if n else ""))
        sug = sugerir_config(colores)
        if sug:
            print("  Sugerencia para config: " + ", ".join(sug) + " (compruébalo mirando una diapositiva).")

    for f in sorted(faltan):
        print(f"\n{f}", file=sys.stderr)
    return 1 if errores else 0


if __name__ == "__main__":
    sys.exit(main())
